# -*- coding: utf-8 -*-
"""
一键流水线: 从 assistant 轨迹 + 质检(evaluator)轨迹 匹配, 筛选, 转 pangu 格式, 输出统计 json。

输入:  assistant 轨迹目录(含 session_report.xlsx) + 质检轨迹目录
输出:  一个输出目录, 内含
         filtered_sessions.txt   筛选保留的 session 名单
         pangu_filtered.jsonl    pangu 格式转换结果(完整对话)
         pangu_filtered_truncated.jsonl / _fold.jsonl  (convert_log 附带产物)
         filter_stats.json       最终统计

步骤:
  1. 读 <assistant_dir>/session_report.xlsx, 保留「错误备注」为空 或 只有"200空响应"的 Session
     (Session 列 == assistant 轨迹的会话子目录名)。
  2. 写 filtered_sessions.txt, 调 convert_log.py --session-list 把筛选后轨迹转 pangu JSONL。
  3. 用「首轮 user 内容(去空白)」把每个筛选后 session 关联到质检轨迹的【Origin Query】,
     取『最早有裁决』的一条作为该 session 的首轮评测结果(completion)。
  4. 输出 filter_stats.json:
       filtered_count       筛选后轨迹数
       with_eval_count      有 evaluator 首轮评测结果的数(含 completion=null 的裁决)
       completion_ge_0.5    completion 为数值且 >=0.5 的数
       completion_eq_1      completion == 1 的数

用法:
  python run_pipeline.py <assistant_dir> <qc_dir> <out_dir>
"""
import os, re, sys, json, glob, argparse, subprocess
from collections import defaultdict

HERE    = os.path.dirname(os.path.abspath(__file__))
CONVERT = os.path.join(HERE, "convert_log.py")

# 「错误备注」中允许保留的良性标记(去掉这些+空白后若为空, 则保留该 session)
BENIGN_NOTES = ["200空响应"]
REPORT_NAME  = "session_report.xlsx"
SHEET_NAME   = "Session详情"


# ── 1. 读 xlsx 并筛选 ──────────────────────────────────────────────────────────
def load_sessions_to_keep(xlsx, sheet):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb[sheet] if sheet in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    idx = {h: i for i, h in enumerate(header)}
    if "Session" not in idx or "错误备注" not in idx:
        raise KeyError(f"报表缺少 Session/错误备注 列; 现有列: {list(header)}")
    sc, ec = idx["Session"], idx["错误备注"]

    kept, dropped = [], []
    for r in rows[1:]:
        sess = r[sc]
        if not sess:
            continue
        note_s = "" if r[ec] is None else str(r[ec]).strip()
        residual = note_s
        for b in BENIGN_NOTES:
            residual = residual.replace(b, "")
        residual = re.sub(r"[\s;,，、]+", "", residual)   # 去分隔符与空白
        if residual == "":
            kept.append(str(sess).strip())
        else:
            dropped.append((str(sess).strip(), note_s))
    return kept, dropped


# ── Origin Query 匹配 & 裁决提取 ──────────────────────────────────────────────
def first_user(path):
    d = json.load(open(path, encoding="utf-8"))
    for m in d.get("messages", []):
        if m.get("role") == "user":
            c = m.get("content")
            return c if isinstance(c, str) else json.dumps(c, ensure_ascii=False)
    return None

_END_RE = re.compile(r"【最近|【验收|【前置|【completion")

def qc_origin_query(path):
    c = first_user(path)
    if not isinstance(c, str) or "【Origin Query】" not in c:
        return None
    seg = c.split("【Origin Query】", 1)[1]
    m = _END_RE.search(seg)
    return (seg[:m.start()] if m else seg).strip()

_TS_HDR_RE = re.compile(r"^\[\w+ \d{4}-\d{2}-\d{2} \d{2}:\d{2} GMT[+-]\d+\]\s*")
_SENDER_META_RE = re.compile(r"^Sender \(untrusted metadata\):\s*```json.*?```\s*", re.DOTALL)

def norm(s):
    """去掉 assistant 侧附加的「[时间戳] Sender (untrusted metadata): ```json...```」外壳,
    再去空白比较, 否则与 evaluator 侧【Origin Query】原文永远无法精确匹配。"""
    s = s or ""
    s = _TS_HDR_RE.sub("", s, count=1)
    s = _SENDER_META_RE.sub("", s, count=1)
    return re.sub(r"\s+", "", s).strip()

def trace_id(path):
    return os.path.basename(os.path.dirname(path))

def _texts(content):
    if isinstance(content, str):
        return content
    if isinstance(content, list):
        return " ".join(x.get("text", "") for x in content if isinstance(x, dict))
    return ""

_INCL_RE = re.compile(r'"inclination"\s*:\s*"(accept|reject|uncertain)"')
_COMP_RE = re.compile(r'"completion"\s*:\s*(null|-?[0-9.]+)')

def qc_verdict(path):
    """(has_verdict, completion_or_None): evaluator 是否给出裁决 + completion 数值。"""
    d = json.load(open(path, encoding="utf-8"))
    full = " ".join(_texts(m.get("content")) for m in d.get("messages", []))
    r = d.get("response")
    if isinstance(r, dict):
        full += " " + _texts(r.get("content"))
    if not _INCL_RE.search(full):
        return False, None
    mc = _COMP_RE.search(full)
    comp = None
    if mc and mc.group(1) != "null":
        try:
            comp = float(mc.group(1))
        except ValueError:
            comp = None
    return True, comp

# 与 convert_log.find_latest_json 一致: 取文件名时间戳最新的 json
def latest_json(subdir):
    fs = glob.glob(os.path.join(subdir, "*.json"))
    def ts(p):
        m = re.search(r'(\d{4}-\d{2}-\d{2})[_T](\d{2})[-:](\d{2})[-:](\d{2})', os.path.basename(p))
        return m.group(0) if m else os.path.basename(p)
    return sorted(fs, key=ts)[-1] if fs else None


# ── 主流程 ────────────────────────────────────────────────────────────────────
def main():
    ap = argparse.ArgumentParser(
        description="轨迹匹配→筛选→转 pangu→统计 一键流水线",
        formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("assistant_dir", help="assistant 轨迹目录(含 session_report.xlsx)")
    ap.add_argument("qc_dir",        help="质检(evaluator)轨迹目录")
    ap.add_argument("out_dir",       help="输出目录(自动创建)")
    a = ap.parse_args()

    for label, d in [("assistant_dir", a.assistant_dir), ("qc_dir", a.qc_dir)]:
        if not os.path.isdir(d):
            ap.error(f"{label} 不存在: {d}")

    os.makedirs(a.out_dir, exist_ok=True)
    session_list = os.path.join(a.out_dir, "filtered_sessions.txt")
    pangu_out    = os.path.join(a.out_dir, "pangu_filtered.jsonl")
    stats_out    = os.path.join(a.out_dir, "filter_stats.json")
    xlsx         = os.path.join(a.assistant_dir, REPORT_NAME)
    if not os.path.exists(xlsx):
        sys.exit(f"[error] 报表不存在: {xlsx}")

    # 1) 筛选
    kept, dropped = load_sessions_to_keep(xlsx, SHEET_NAME)
    with open(session_list, "w", encoding="utf-8") as f:
        f.write("\n".join(kept) + ("\n" if kept else ""))
    print(f"[1] 筛选: 保留 {len(kept)}  丢弃 {len(dropped)}")
    for s, n in dropped:
        print(f"      drop {s}  <- {n!r}")

    # 2) 转 pangu(先清旧输出, convert_log 是追加写)
    for f in (pangu_out,
              pangu_out.replace(".jsonl", "_truncated.jsonl"),
              pangu_out.replace(".jsonl", "_fold.jsonl")):
        if os.path.exists(f):
            os.remove(f)
    cmd = [sys.executable, CONVERT, a.assistant_dir, pangu_out, "--session-list", session_list]
    print(f"[2] 转换: convert_log.py -> {os.path.basename(pangu_out)}")
    res = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8")
    for line in (res.stdout or "").splitlines()[-2:]:
        print("      ", line)
    if res.returncode != 0:
        print("      [stderr]", (res.stderr or "")[-500:])
    n_complete  = sum(1 for _ in open(pangu_out, encoding="utf-8")) if os.path.exists(pangu_out) else 0
    trunc_path  = pangu_out.replace(".jsonl", "_truncated.jsonl")
    n_truncated = sum(1 for _ in open(trunc_path, encoding="utf-8")) if os.path.exists(trunc_path) else 0

    # 3) QC 按归一化 Origin Query 分组
    qfiles = sorted(glob.glob(os.path.join(a.qc_dir, "*", "*.json")))
    groups = defaultdict(list)
    for p in qfiles:
        groups[norm(qc_origin_query(p))].append(p)

    def first_eval_for(qkey):
        verds = []
        for p in groups.get(qkey, []):
            hv, comp = qc_verdict(p)
            if hv:
                verds.append((trace_id(p), comp))
        if not verds:
            return None
        verds.sort(key=lambda x: x[0])   # 目录名时间戳升序 -> 最早
        return verds[0]

    # 4) 关联评测并统计
    rows = []
    with_eval = ge05 = eq1 = 0
    for sess in kept:
        p = latest_json(os.path.join(a.assistant_dir, sess))
        qkey = norm(first_user(p)) if p else None
        ev = first_eval_for(qkey) if qkey else None
        if ev:
            with_eval += 1
            comp = ev[1]
            if comp is not None and comp >= 0.5:
                ge05 += 1
            if comp is not None and comp == 1:
                eq1 += 1
        rows.append({"session": sess, "has_eval": bool(ev),
                     "eval_qc": ev[0] if ev else "",
                     "completion": ev[1] if ev else None})

    stats = {
        "filtered_count":    len(kept),
        "with_eval_count":   with_eval,
        "completion_ge_0.5": ge05,
        "completion_eq_1":   eq1,
        "note": "with_eval_count 含 completion=null 的裁决; ge_0.5/eq_1 仅统计 completion 为数值者",
        "pangu_complete":    n_complete,
        "pangu_truncated":   n_truncated,
        "dropped_count":     len(dropped),
        "per_session":       rows,
    }
    with open(stats_out, "w", encoding="utf-8") as f:
        json.dump(stats, f, ensure_ascii=False, indent=2)

    print(f"[3] pangu: complete {n_complete} 条, truncated {n_truncated} 条")
    print(f"[4] 统计 -> {stats_out}")
    print(f"      筛选后轨迹数           : {len(kept)}")
    print(f"      有 evaluator 首轮评测   : {with_eval}")
    print(f"      completion >= 0.5      : {ge05}")
    print(f"      completion == 1        : {eq1}")


if __name__ == "__main__":
    main()
